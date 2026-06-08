"""Knowledge Service - 知识文档管理"""
import json
import os
import uuid
import re
from datetime import datetime
import sys

backend_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from backend.core.database import get_db, execute_query, execute_insert
from backend.services.rag_engine import init_chroma, add_documents, search_documents

def generate_item_id():
    return f"KNW-{uuid.uuid4().hex[:12]}"

def process_uploaded_file(file_content: bytes, file_name: str, doc_type: str, description: str = "") -> dict:
    item_id = generate_item_id()
    text_content = extract_text(file_content, file_name)
    chunks = split_text(text_content)
    if chunks:
        metadatas = [{"type": doc_type, "file": file_name, "chunk": i} for i in range(len(chunks))]
        ids = [f"{item_id}_{i}" for i in range(len(chunks))]
        add_documents(chunks, metadatas, ids)
    try:
        with get_db() as conn:
            execute_insert(conn, """
                INSERT INTO knowledge_items (item_id, type, file_name, description, status, progress, message)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """, (item_id, doc_type, file_name, description, "completed", 1.0,
                  f"成功处理 {len(chunks)} 个文本块"))
    except Exception as e:
        print(f"DB save knowledge error: {e}")
    return {"task_id": item_id, "status": "completed", "file_name": file_name}

def extract_text(file_content: bytes, file_name: str) -> str:
    ext = file_name.lower().split('.')[-1]
    if ext in ('txt', 'md'):
        return file_content.decode('utf-8', errors='ignore')
    elif ext == 'docx':
        try:
            from docx import Document
            from io import BytesIO
            doc = Document(BytesIO(file_content))
            return '\n'.join([p.text for p in doc.paragraphs])
        except:
            return ""
    elif ext == 'xlsx':
        try:
            import openpyxl
            from io import BytesIO
            wb = openpyxl.load_workbook(BytesIO(file_content))
            texts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    texts.append(' '.join(str(c) for c in row if c))
            return '\n'.join(texts)
        except:
            return ""
    else:
        return file_content.decode('utf-8', errors='ignore')

def split_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list:
    if not text:
        return []
    paragraphs = re.split(r'\n\s*\n', text)
    chunks = []
    current_chunk = ""
    for para in paragraphs:
        if len(current_chunk) + len(para) > chunk_size:
            if current_chunk:
                chunks.append(current_chunk.strip())
            current_chunk = para
        else:
            current_chunk += ("\n" + para if current_chunk else para)
    if current_chunk:
        chunks.append(current_chunk.strip())
    return chunks[:100]